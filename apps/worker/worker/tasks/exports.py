from __future__ import annotations

import asyncio
import io
import os
import uuid
from datetime import UTC, datetime
from html import escape
from pathlib import Path

from sqlalchemy import select

from app.core.auth import load_principal
from app.core.config import get_settings
from app.core.resources import Resources
from app.core.storage import LocalStorage
from app.db.tenant import system_session, tenant_session
from app.domain.tabular import safe_cell
from app.models import ExportJob, Tenant
from app.models.enums import ExportKind, JobStatus
from app.schemas.feedback import FeedbackFilter
from app.services.analytics import AnalyticsService
from worker.celery_app import celery_app


async def render(service: AnalyticsService, filters: FeedbackFilter, kind: ExportKind, app_name: str) -> bytes:
    summary = await service.overview(filters)
    topics = await service.topics(filters)
    output = io.BytesIO()
    if kind == ExportKind.XLSX:
        from openpyxl import Workbook
        workbook = Workbook(write_only=True)
        overview = workbook.create_sheet("Tóm tắt"); overview.append([app_name]); overview.append(["Từ", str(filters.start or ""), "Đến trước", str(filters.end or "")])
        overview.append(["Tổng phản hồi", summary["total_responses"]]); overview.append(["CSAT", summary["csat"]]); overview.append(["Khẩn cấp", summary["urgent_responses"]])
        for label, value in summary["sentiments"].items():
            overview.append([{"positive": "Tích cực", "negative": "Tiêu cực", "neutral": "Trung lập"}[label], value["count"], value["percent"]])
        detail = workbook.create_sheet("Chi tiết")
        detail.append(["Nội dung", "Sao", "Cảm xúc", "Chủ đề", "Khẩn cấp", "Nguồn", "Thời gian"])
        query = await service.analysis_scope(filters)
        result = await service.db.stream_scalars(query.execution_options(yield_per=200))
        async for row in result:
            detail.append([safe_cell(row.text), float(row.rating) if row.rating is not None else None, row.sentiment, ",".join(str(t) for t in row.topic_ids), row.is_urgent, row.channel, row.responded_at.isoformat()])
        sheet = workbook.create_sheet("Theo chủ đề"); sheet.append(["Kỳ", "Mã chủ đề", "Số nhận xét"])
        for topic in topics:
            sheet.append([topic["date"], topic["topic_id"], topic["count"]])
        workbook.save(output)
    else:
        from reportlab.graphics.charts.barcharts import VerticalBarChart
        from reportlab.graphics.shapes import Drawing
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table
        font = os.environ.get("REPORT_FONT_PATH", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
        if not Path(font).is_file():
            raise RuntimeError("Thiếu font Unicode; đặt REPORT_FONT_PATH để xuất PDF tiếng Việt.")
        pdfmetrics.registerFont(TTFont("ReportUnicode", font))
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = "ReportUnicode"
        content = [Paragraph(escape(app_name), styles["Title"]), Paragraph("Báo cáo phản hồi khách hàng", styles["Heading1"]), Paragraph(escape(f"Từ {filters.start or 'đầu kỳ'} đến trước {filters.end or 'hiện tại'}"), styles["Normal"]), Spacer(1, 16)]
        table = Table([["Tổng phản hồi", str(summary["total_responses"])], ["CSAT", str(summary["csat"] or "—")], ["Cảnh báo khẩn", str(summary["urgent_responses"])]])
        table.setStyle([("FONTNAME", (0,0), (-1,-1), "ReportUnicode"), ("GRID", (0,0),(-1,-1),0.5,colors.lightgrey)])
        content.append(table)
        chart = VerticalBarChart(); chart.x=40; chart.y=30; chart.width=360; chart.height=170
        chart.data = [[summary["sentiments"][s]["count"] for s in ("positive","negative","neutral")]]
        chart.categoryAxis.categoryNames=["Tích cực","Tiêu cực","Trung lập"]; chart.categoryAxis.labels.fontName="ReportUnicode"
        chart.bars[(0,0)].fillColor=colors.green; chart.bars[(0,1)].fillColor=colors.red; chart.bars[(0,2)].fillColor=colors.grey
        drawing=Drawing(440,230); drawing.add(chart); content.append(drawing)
        content.append(Paragraph("Tỷ lệ cảm xúc tính trên nhận xét văn bản đã phân tích. Dữ liệu định danh khách không được đưa vào báo cáo.", styles["Normal"]))
        SimpleDocTemplate(output).build(content)
    return output.getvalue()


async def run(tenant_id: str, job_id: str) -> None:
    resources = Resources.from_settings(get_settings())
    try:
        async with tenant_session(resources.session_factory, uuid.UUID(tenant_id)) as db:
            job = (await db.execute(select(ExportJob).where(ExportJob.id == uuid.UUID(job_id), ExportJob.tenant_id == uuid.UUID(tenant_id)).with_for_update(skip_locked=True))).scalar_one_or_none()
            if job is None or job.status == JobStatus.DONE:
                return
            job.status = JobStatus.RUNNING
            try:
                if job.created_by is None:
                    raise RuntimeError("Người tạo tác vụ không còn tồn tại.")
                principal = await load_principal(db, job.created_by, job.tenant_id, None)
                from app.core.permissions import Permission
                principal.require(Permission.DATA_EXPORT)
                filters = FeedbackFilter.model_validate(job.params["filters"])
                content = await render(AnalyticsService(db, principal), filters, job.kind, get_settings().app_name)
                suffix = "pdf" if job.kind == ExportKind.PDF else "xlsx"
                key = f"{tenant_id}/exports/{job_id}.{suffix}"
                await asyncio.to_thread(LocalStorage(get_settings().storage_local_root).put, key, content)
                job.file_key = key; job.filename = f"bao-cao-{job_id[:8]}.{suffix}"; job.status = JobStatus.DONE; job.finished_at=datetime.now(UTC)
            except Exception as exc:
                job.status = JobStatus.FAILED; job.error_message = type(exc).__name__
    finally:
        await resources.close()


@celery_app.task(name="worker.tasks.exports.run")
def export_file(tenant_id: str, job_id: str) -> None:
    asyncio.run(run(tenant_id, job_id))


async def scan() -> int:
    resources=Resources.from_settings(get_settings()); count=0
    try:
        async with system_session(resources.session_factory) as db:
            tenants=list((await db.execute(select(Tenant.id).where(Tenant.deleted_at.is_(None)))).scalars())
        for tid in tenants:
            async with tenant_session(resources.session_factory,tid) as db:
                jobs=list((await db.execute(select(ExportJob.id).where(ExportJob.tenant_id==tid,ExportJob.status.in_([JobStatus.PENDING,JobStatus.RUNNING])))).scalars())
            for jid in jobs:
                resources.queue.send("worker.tasks.exports.run",kwargs={"tenant_id":str(tid),"job_id":str(jid)});count+=1
        return count
    finally:
        await resources.close()


@celery_app.task(name="worker.tasks.exports.scan")
def scan_exports() -> int:
    return asyncio.run(scan())
